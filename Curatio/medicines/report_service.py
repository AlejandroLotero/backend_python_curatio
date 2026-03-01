"""
Servicio para generar reportes de medicamentos en Excel y PDF
RFADMIN13: Generar reportes de medicamentos
"""

import io
from datetime import datetime
from decimal import Decimal
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.pdfgen import canvas
from .models import Medicine


class MedicineReportService:
    """Servicio centralizado para generar reportes de medicamentos"""
    
    REPORT_FIELDS = [
        ('name', 'Nombre'),
        ('pharmaceutical_form', 'Forma Farmacéutica'),
        ('presentation', 'Presentación'),
        ('concentration', 'Concentración'),
        ('administration_route', 'Vía de Administración'),
        ('laboratory', 'Laboratorio'),
        ('batch_number', 'Lote'),
        ('manufacture_date', 'Fecha de Fabricación'),
        ('expiration_date', 'Fecha de Vencimiento'),
        ('stock_quantity', 'Stock'),
        ('purchase_price', 'Precio de Compra'),
        ('sale_price', 'Precio de Venta'),
        ('supplier', 'Proveedor'),
        ('requires_formula', 'Requiere Fórmula'),
        ('description', 'Descripción'),
        ('state', 'Estado'),
    ]
    
    @staticmethod
    def filter_medicines_by_state(state=None):
        """
        Filtra medicamentos según su estado
        
        Args:
            state: str ('ACTIVO', 'VENCIDO', 'AGOTADO', 'SUSPENDIDO', o None para todos)
        
        Returns:
            QuerySet de medicamentos filtrados
        """
        medicines = Medicine.objects.all().order_by('name')
        
        if state and state != 'TODOS':
            if state not in ['ACTIVO', 'VENCIDO', 'AGOTADO', 'SUSPENDIDO']:
                raise ValueError(f"Estado inválido: {state}")
            medicines = medicines.filter(state=state)
        
        return medicines
    
    @staticmethod
    def prepare_report_data(medicines):
        """
        Prepara los datos de medicamentos para el reporte
        
        Args:
            medicines: QuerySet de medicamentos
        
        Returns:
            Lista de diccionarios con datos formateados
        """
        report_data = []
        
        for medicine in medicines:
            row = {
                'name': medicine.name,
                'pharmaceutical_form': medicine.get_pharmaceutical_form_display(),
                'presentation': medicine.presentation,
                'concentration': medicine.concentration,
                'administration_route': medicine.get_administration_route_display(),
                'laboratory': medicine.get_laboratory_display(),
                'batch_number': medicine.batch_number,
                'manufacture_date': medicine.manufacture_date.strftime('%Y-%m-%d'),
                'expiration_date': medicine.expiration_date.strftime('%Y-%m-%d'),
                'stock_quantity': medicine.stock_quantity,
                'purchase_price': f"${medicine.purchase_price:,.2f}",
                'sale_price': f"${medicine.sale_price:,.2f}",
                'supplier': medicine.supplier,
                'requires_formula': 'Sí' if medicine.requires_formula else 'No',
                'description': medicine.description or '-',
                'state': medicine.get_state_display(),
            }
            report_data.append(row)
        
        return report_data
    
    @staticmethod
    def get_metadata(state=None, user=None):
        """
        Genera metadatos del reporte
        
        Args:
            state: Estado filtrado
            user: Usuario que generó el reporte
        
        Returns:
            Diccionario con metadatos
        """
        full_name = 'Sistema'
        if user:
            first_name = getattr(user, 'first_name', '')
            last_name = getattr(user, 'last_name', '')
            full_name = f"{first_name} {last_name}".strip() or user.username or 'Sistema'
        
        return {
            'generated_at': timezone.now(),
            'generated_at_str': timezone.now().strftime('%d/%m/%Y %H:%M:%S'),
            'generated_by': full_name,
            'state_filter': state or 'TODOS',
            'system_name': 'SISTEMA DE GESTIÓN FARMACÉUTICA - CURATIO',
        }


class ExcelReportGenerator:
    """Generador de reportes en formato Excel"""
    
    SERVICE = MedicineReportService()
    
    @classmethod
    def generate(cls, state=None, user=None):
        """
        Genera un reporte en Excel
        
        Args:
            state: Estado a filtrar
            user: Usuario que solicita el reporte
        
        Returns:
            HttpResponse con archivo Excel
        """
        # Obtener datos
        medicines = cls.SERVICE.filter_medicines_by_state(state)
        
        if not medicines.exists():
            raise ValueError(f"No hay medicamentos con el estado: {state or 'TODOS'}")
        
        report_data = cls.SERVICE.prepare_report_data(medicines)
        metadata = cls.SERVICE.get_metadata(state, user)
        
        # Crear workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Reporte Medicamentos"
        
        # === ENCABEZADO ===
        cls._add_header(ws, metadata)
        
        # === TABLA DE DATOS ===
        current_row = cls._add_data_table(ws, report_data)
        
        # === RESUMEN ===
        cls._add_summary(ws, report_data, current_row)
        
        # Ajustar ancho de columnas
        cls._adjust_columns(ws)
        
        # Crear respuesta HTTP
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        filename = f"Reporte_Medicamentos_{state or 'TODOS'}_{metadata['generated_at'].strftime('%Y%m%d_%H%M%S')}.xlsx"
        
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        return response
    
    @staticmethod
    def _add_header(ws, metadata):
        """Agrega el encabezado al reporte"""
        # Título del sistema
        ws['A1'] = metadata['system_name']
        ws['A1'].font = Font(size=14, bold=True, color="FFFFFF")
        ws['A1'].fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        ws.merge_cells('A1:P1')
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 25
        
        # Información de generación
        ws['A2'] = f"Fecha de Generación: {metadata['generated_at_str']}"
        ws['A2'].font = Font(size=10, italic=True)
        
        ws['A3'] = f"Generado por: {metadata['generated_by']}"
        ws['A3'].font = Font(size=10, italic=True)
        
        ws['A4'] = f"Filtro de Estado: {metadata['state_filter']}"
        ws['A4'].font = Font(size=10, italic=True)
        
        # Línea en blanco
        ws.row_dimensions[5].height = 5
    
    @staticmethod
    def _add_data_table(ws, report_data):
        """Agrega la tabla de datos al reporte"""
        # Encabezados de columna
        row_num = 6
        headers = [field[1] for field in MedicineReportService.REPORT_FIELDS]
        field_keys = [field[0] for field in MedicineReportService.REPORT_FIELDS]
        
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        ws.row_dimensions[row_num].height = 30
        
        # Datos
        row_num += 1
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for data_row in report_data:
            for col_num, key in enumerate(field_keys, 1):
                cell = ws.cell(row=row_num, column=col_num)
                cell.value = data_row.get(key, '-')
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
                
                # Colorear según estado
                if key == 'state':
                    state_value = data_row.get('state', '')
                    if state_value == 'Activo':
                        cell.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
                    elif state_value == 'Vencido':
                        cell.fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
                    elif state_value == 'Agotado':
                        cell.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
                    elif state_value == 'Suspendido':
                        cell.fill = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")
            
            ws.row_dimensions[row_num].height = 20
            row_num += 1
        
        return row_num
    
    @staticmethod
    def _add_summary(ws, report_data, start_row):
        """Agrega un resumen del reporte"""
        ws.row_dimensions[start_row].height = 10
        
        summary_row = start_row + 1
        
        # Conteos por estado
        states_count = {}
        for row in report_data:
            state = row['state']
            states_count[state] = states_count.get(state, 0) + 1
        
        ws[f'A{summary_row}'] = "RESUMEN"
        ws[f'A{summary_row}'].font = Font(bold=True, size=11)
        summary_row += 1
        
        ws[f'A{summary_row}'] = "Total de Medicamentos:"
        ws[f'B{summary_row}'] = len(report_data)
        ws[f'A{summary_row}'].font = Font(bold=True)
        summary_row += 1
        
        for state, count in sorted(states_count.items()):
            ws[f'A{summary_row}'] = f"  • {state}:"
            ws[f'B{summary_row}'] = count
            summary_row += 1
    
    @staticmethod
    def _adjust_columns(ws):
        """Ajusta el ancho de las columnas"""
        column_widths = [20, 18, 22, 15, 18, 25, 12, 18, 18, 10, 15, 15, 20, 15, 25, 12]
        
        for col_num, width in enumerate(column_widths, 1):
            ws.column_dimensions[chr(64 + col_num)].width = width


class PDFReportGenerator:
    """Generador de reportes en formato PDF"""
    
    SERVICE = MedicineReportService()
    
    @classmethod
    def generate(cls, state=None, user=None):
        """
        Genera un reporte en PDF
        
        Args:
            state: Estado a filtrar
            user: Usuario que solicita el reporte
        
        Returns:
            HttpResponse con archivo PDF
        """
        # Obtener datos
        medicines = cls.SERVICE.filter_medicines_by_state(state)
        
        if not medicines.exists():
            raise ValueError(f"No hay medicamentos con el estado: {state or 'TODOS'}")
        
        report_data = cls.SERVICE.prepare_report_data(medicines)
        metadata = cls.SERVICE.get_metadata(state, user)
        
        # Crear PDF
        output = io.BytesIO()
        doc = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=20,
            leftMargin=20,
            topMargin=30,
            bottomMargin=30,
        )
        
        # Elementos del documento
        elements = []
        
        # Encabezado
        elements.extend(cls._create_header(metadata))
        
        # Tabla de datos
        elements.extend(cls._create_data_table(report_data))
        
        # Resumen
        elements.extend(cls._create_summary(report_data))
        
        # Construir PDF
        doc.build(elements)
        output.seek(0)
        
        filename = f"Reporte_Medicamentos_{state or 'TODOS'}_{metadata['generated_at'].strftime('%Y%m%d_%H%M%S')}.pdf"
        
        response = HttpResponse(
            output.getvalue(),
            content_type='application/pdf'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        
        return response
    
    @staticmethod
    def _create_header(metadata):
        """Crea el encabezado del PDF"""
        styles = getSampleStyleSheet()
        elements = []
        
        # Título
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=16,
            textColor=colors.HexColor('#1F4E78'),
            spaceAfter=6,
            alignment=1,  # center
            fontName='Helvetica-Bold'
        )
        
        elements.append(Paragraph(metadata['system_name'], title_style))
        
        # Información de generación
        info_style = ParagraphStyle(
            'Info',
            parent=styles['Normal'],
            fontSize=9,
            textColor=colors.grey,
            spaceAfter=3,
        )
        
        elements.append(Paragraph(f"Fecha de Generación: {metadata['generated_at_str']}", info_style))
        elements.append(Paragraph(f"Generado por: {metadata['generated_by']}", info_style))
        elements.append(Paragraph(f"Filtro de Estado: {metadata['state_filter']}", info_style))
        elements.append(Spacer(1, 12))
        
        return elements
    
    @staticmethod
    def _create_data_table(report_data):
        """Crea la tabla de datos del PDF"""
        elements = []
        
        # Preparar datos de la tabla
        headers = [field[1] for field in MedicineReportService.REPORT_FIELDS]
        field_keys = [field[0] for field in MedicineReportService.REPORT_FIELDS]
        
        table_data = [headers]
        for data_row in report_data:
            row = [str(data_row.get(key, '-'))for key in field_keys]
            table_data.append(row)
        
        # Crear tabla
        table = Table(table_data, colWidths=[
            0.9*inch, 1.0*inch, 1.1*inch, 0.8*inch, 1.0*inch, 1.2*inch,
            0.7*inch, 0.9*inch, 0.9*inch, 0.6*inch, 0.9*inch, 0.9*inch,
            1.2*inch, 0.8*inch, 1.2*inch, 0.7*inch
        ])
        
        # Estilos de tabla
        table_style = TableStyle([
            # Encabezado
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4472C4')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            
            # Datos
            ('ALIGN', (0, 1), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 7),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F2F2F2')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ])
        
        table.setStyle(table_style)
        elements.append(table)
        elements.append(Spacer(1, 12))
        
        return elements
    
    @staticmethod
    def _create_summary(report_data):
        """Crea el resumen del PDF"""
        elements = []
        styles = getSampleStyleSheet()
        
        # Conteos por estado
        states_count = {}
        for row in report_data:
            state = row['state']
            states_count[state] = states_count.get(state, 0) + 1
        
        summary_style = ParagraphStyle(
            'Summary',
            parent=styles['Normal'],
            fontSize=9,
            spaceAfter=4,
        )
        
        elements.append(Paragraph("<b>RESUMEN</b>", summary_style))
        elements.append(Paragraph(f"Total de Medicamentos: <b>{len(report_data)}</b>", summary_style))
        
        for state, count in sorted(states_count.items()):
            elements.append(Paragraph(f"• {state}: <b>{count}</b>", summary_style))
        
        return elements
